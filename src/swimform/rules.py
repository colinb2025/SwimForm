"""
Load swim_rules.xlsx, validate it, and evaluate measurements against it.

Usage:
  python -m swimform.rules rules/swim_rules.xlsx            # validate + export JSON
  python -m swimform.rules rules/swim_rules.xlsx --demo     # also run a sample analysis
"""
import json, sys
from collections import defaultdict
from openpyxl import load_workbook

ENUMS = {
    "view": {"side", "head_on"},
    "water": {"above", "below", "either"},
    "operator": {"<", "<=", ">", ">=", "between", "outside"},
    "type": {"drill", "exercise", "cue"},
    "status": {"draft", "verified", "disabled"},
}

def read_sheet(wb, name):
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [h for h in rows[0]]
    out = []
    for r in rows[1:]:
        if r[0] in (None, ""):
            continue
        out.append({h: v for h, v in zip(hdr, r) if h})
    return out

def load(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    rules = {
        "metrics": read_sheet(wb, "Metrics"),
        "issues": read_sheet(wb, "Issues"),
        "issue_rules": read_sheet(wb, "Issue_Rules"),
        "interventions": read_sheet(wb, "Interventions"),
        "issue_interventions": read_sheet(wb, "Issue_Interventions"),
    }
    errors = validate(rules)
    # drop disabled rows after validation
    for k in ("metrics", "issues", "interventions"):
        rules[k] = [r for r in rules[k] if r.get("status") != "disabled"]
    return rules, errors

def validate(R):
    errs = []
    def uniq(rows, key, sheet):
        seen = set()
        for r in rows:
            if r[key] in seen: errs.append(f"{sheet}: duplicate {key} '{r[key]}'")
            seen.add(r[key])
        return seen
    m_ids = uniq(R["metrics"], "metric_id", "Metrics")
    i_ids = uniq(R["issues"], "issue_id", "Issues")
    v_ids = uniq(R["interventions"], "intervention_id", "Interventions")
    for m in R["metrics"]:
        for f in ("view", "water", "status"):
            if m.get(f) not in ENUMS[f]: errs.append(f"Metrics/{m['metric_id']}: bad {f} '{m.get(f)}'")
        lo, hi = m.get("ideal_min"), m.get("ideal_max")
        if lo is not None and hi is not None and lo > hi:
            errs.append(f"Metrics/{m['metric_id']}: ideal_min > ideal_max")
        if m.get("fault_low") is not None and lo is not None and m["fault_low"] >= lo:
            errs.append(f"Metrics/{m['metric_id']}: fault_low must be below ideal_min")
        if m.get("fault_high") is not None and hi is not None and m["fault_high"] <= hi:
            errs.append(f"Metrics/{m['metric_id']}: fault_high must be above ideal_max")
    for r in R["issue_rules"]:
        tag = f"Issue_Rules/{r['issue_id']}"
        if r["issue_id"] not in i_ids: errs.append(f"{tag}: unknown issue_id")
        if r["metric_id"] not in m_ids: errs.append(f"{tag}: unknown metric_id '{r['metric_id']}'")
        if r["operator"] not in ENUMS["operator"]: errs.append(f"{tag}: bad operator '{r['operator']}'")
        if r["operator"] in ("between", "outside") and r.get("value2") is None:
            errs.append(f"{tag}: '{r['operator']}' needs value2")
    for v in R["interventions"]:
        if v.get("type") not in ENUMS["type"]: errs.append(f"Interventions/{v['intervention_id']}: bad type")
    for r in R["issue_interventions"]:
        if r["issue_id"] not in i_ids: errs.append(f"Issue_Interventions: unknown issue '{r['issue_id']}'")
        if r["intervention_id"] not in v_ids: errs.append(f"Issue_Interventions: unknown intervention '{r['intervention_id']}'")
    ruled = {r["issue_id"] for r in R["issue_rules"]}
    for i in i_ids - ruled: errs.append(f"Issues/{i}: no rules, can never trigger")
    return errs

def score(m, x):
    """100 inside ideal range, linear to 0 at the fault bound, 0 beyond it."""
    lo, hi, flo, fhi = m.get("ideal_min"), m.get("ideal_max"), m.get("fault_low"), m.get("fault_high")
    if lo is not None and x < lo:
        if flo is None: return 100.0 if hi is None else max(0.0, 100 * (1 - (lo - x) / max(abs(lo), 1)))
        return max(0.0, 100 * (x - flo) / (lo - flo))
    if hi is not None and x > hi:
        if fhi is None: return max(0.0, 100 * (1 - (x - hi) / max(abs(hi), 1)))
        return max(0.0, 100 * (fhi - x) / (fhi - hi))
    return 100.0

def check(op, x, a, b=None):
    return {"<": x < a, "<=": x <= a, ">": x > a, ">=": x >= a,
            "between": b is not None and a <= x <= b,
            "outside": b is not None and (x < a or x > b)}[op]

def overshoot(op, x, a, b=None):
    """How far past the threshold, relative to its size (for severity)."""
    ref = max(abs(a), 1e-6)
    if op in ("<", "<="): return (a - x) / ref
    if op in (">", ">="): return (x - a) / ref
    return 0.1

def evaluate(R, measurements, per_cycle=None, top_n=3):
    """
    measurements: {metric_id: aggregated value}
    per_cycle:    optional {metric_id: [value per usable cycle]} for min_cycle_pct rules
    """
    per_cycle = per_cycle or {}
    metrics = {m["metric_id"]: m for m in R["metrics"]}
    scores = {mid: round(score(metrics[mid], v)) for mid, v in measurements.items()
              if mid in metrics and metrics[mid].get("scored") == "yes"}

    groups = defaultdict(lambda: defaultdict(list))
    for r in R["issue_rules"]:
        groups[r["issue_id"]][r.get("rule_group") or 1].append(r)

    issues = {i["issue_id"]: i for i in R["issues"]}
    found = []
    for iid, grp in groups.items():
        if iid not in issues: continue
        best = None
        for rows in grp.values():
            ok, sev = True, 0.0
            for r in rows:
                mid = r["metric_id"]
                if mid not in measurements: ok = False; break
                pct = r.get("min_cycle_pct")
                if pct is not None and per_cycle.get(mid):
                    vals = per_cycle[mid]
                    share = sum(check(r["operator"], v, r["value"], r.get("value2")) for v in vals) / len(vals)
                    if share < pct: ok = False; break
                elif not check(r["operator"], measurements[mid], r["value"], r.get("value2")):
                    ok = False; break
                sev = max(sev, overshoot(r["operator"], measurements[mid], r["value"], r.get("value2")))
            if ok: best = max(best or 0, sev)
        if best is not None:
            w = issues[iid]["priority_weight"]
            found.append({"issue_id": iid, "name": issues[iid]["name"], "priority_weight": w,
                          "severity": round(best, 3), "priority": round(w * (1 + best), 3)})
    found.sort(key=lambda f: f["priority"], reverse=True)

    ints = {v["intervention_id"]: v for v in R["interventions"]}
    by_issue = defaultdict(list)
    for r in sorted(R["issue_interventions"], key=lambda r: r["rank"]):
        if r["intervention_id"] in ints: by_issue[r["issue_id"]].append(ints[r["intervention_id"]])
    for f in found[:top_n]:
        i = issues[f["issue_id"]]
        f.update(what_we_observed=i["what_we_observed"], why_it_matters=i["why_it_matters"],
                 interventions=[{"id": v["intervention_id"], "type": v["type"], "name": v["name"],
                                 "equipment": v.get("equipment") or ""} for v in by_issue[f["issue_id"]]])
    return {"scores": scores, "priorities": found[:top_n], "other_issues": [f["issue_id"] for f in found[top_n:]]}

if __name__ == "__main__":
    path = sys.argv[1] if len(sys.argv) > 1 else "swim_rules.xlsx"
    R, errors = load(path)
    if errors:
        print("VALIDATION ERRORS:"); [print("  -", e) for e in errors]; sys.exit(1)
    print(f"OK: {len(R['metrics'])} metrics, {len(R['issues'])} issues, "
          f"{len(R['issue_rules'])} rules, {len(R['interventions'])} interventions")
    with open(path.rsplit(".", 1)[0] + ".json", "w") as f:
        json.dump(R, f, indent=2, default=str)
    if "--demo" in sys.argv:
        sample = {"catch_upper_arm_angle": 52, "catch_forearm_angle": 38, "mid_pull_elbow_angle": 118,
                  "head_lift_angle": 26, "body_line_angle": 12, "shoulder_rotation": 38,
                  "entry_crossover": 0.35, "kick_knee_angle": 132, "stroke_rate": 30}
        cycles = {"entry_crossover": [0.4, 0.33, 0.1, 0.38]}
        print(json.dumps(evaluate(R, sample, cycles), indent=2))
